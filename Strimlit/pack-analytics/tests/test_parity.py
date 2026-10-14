"""Порівняння з незмінним оригіналом, а не повторення формул у тестах."""
import ast
from io import BytesIO
from pathlib import Path
import contextlib
import io
import pandas as pd
import pytest
from openpyxl import load_workbook
from packing.calculations import calculate
from packing.demo import demo_frames, as_excel
from packing.excel import export_excel
from packing.inputs import read_inputs

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize('night,extended', [(False, False), (False, True), (True, False), (True, True)])
def test_all_tables_and_excel_match_original(night, extended, tmp_path, monkeypatch):
    pack, employees = demo_frames(night, extended)
    pack.to_excel(tmp_path/'Pack_data_06.10.2026.xlsx', index=False)
    employees.to_excel(tmp_path/'InfUser.xlsx', index=False)
    monkeypatch.chdir(tmp_path)
    scope = {'__name__': '__main__'}
    with contextlib.redirect_stdout(io.StringIO()):
        exec(compile((ROOT/'reference/original.py').read_text(), 'original.py', 'exec'), scope)
    # Compare using the same Excel-decoded dtypes the original script receives.
    pack, employees = read_inputs(as_excel(pack), as_excel(employees))
    result = calculate(pack, employees)
    for name, frame in result['tables'].items():
        pd.testing.assert_frame_equal(frame, scope[name], check_exact=True)
    assert result['night'] == night
    assert result['extended'] == extended
    original = load_workbook(tmp_path/'Pack_report.xlsx')
    migrated = load_workbook(BytesIO(export_excel(result['tables'])))
    assert original.sheetnames == migrated.sheetnames
    for sheet in original.sheetnames:
        assert list(original[sheet].values) == list(migrated[sheet].values)
        for row_a, row_b in zip(original[sheet], migrated[sheet]):
            for a,b in zip(row_a,row_b):
                assert a.style_id == b.style_id


def test_calculation_statements_are_unchanged():
    original = ast.parse((ROOT/'reference/original.py').read_text())
    migrated = ast.parse((ROOT/'packing/calculations.py').read_text())
    helpers = {'get_status','find_break','cal_rate'}
    calc = next(n for n in migrated.body if isinstance(n,ast.FunctionDef) and n.name == 'calculate')
    all_new = migrated.body + calc.body
    for name in helpers:
        before = next(n for n in original.body if isinstance(n,ast.FunctionDef) and n.name == name)
        after = next(n for n in all_new if isinstance(n,ast.FunctionDef) and n.name == name)
        assert ast.dump(before) == ast.dump(after)
    # All calculation statements between input preparation and Excel export are identical.
    source = (ROOT/'reference/original.py').read_text()
    body = ast.parse(source[source.index('data_employee["Create time"]='):source.index('# 20. ЕКСПОРТ В EXCEL')]).body
    body = [n for n in body if not (isinstance(n,ast.Expr) and isinstance(n.value,ast.Call)
            and isinstance(n.value.func,ast.Name) and n.value.func.id == 'print')]
    assert [ast.dump(n) for n in body] == [ast.dump(n) for n in calc.body[3:-1]]


def test_input_validation():
    pack, employees = demo_frames()
    with pytest.raises(ValueError, match='відсутні колонки'):
        read_inputs(as_excel(pack.drop(columns='Operator')), as_excel(employees))
    with pytest.raises(ValueError, match='не вдалося прочитати'):
        read_inputs(b'invalid workbook', as_excel(employees))
