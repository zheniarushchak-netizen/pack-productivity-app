"""Оригінальне оформлення та розташування таблиць Excel."""
from io import BytesIO
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment

def write_table(
    writer,
    dataframe,
    sheet_name,
    startrow=0,
    startcol=0,
    index=True,
    header=True,
    format_headers=True
):
    dataframe.to_excel(
        writer,
        sheet_name=sheet_name,
        startrow=startrow,
        startcol=startcol,
        index=index,
        header=header
    )

    if format_headers:
        worksheet = writer.book[sheet_name]

        format_table_headers(
            worksheet,
            dataframe,
            startrow=startrow,
            startcol=startcol,
            index=index
        )

def format_table_headers(
    worksheet,
    dataframe,
    startrow=0,
    startcol=0,
    index=True,
    fill_color="D9EAF7"
):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor=fill_color
    )

    header_font = Font(
        bold=True
    )

    header_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    excel_start_row = startrow + 1
    excel_start_col = startcol + 1

    header_rows = dataframe.columns.nlevels

    if index:
        index_cols = dataframe.index.nlevels
    else:
        index_cols = 0

    table_width = len(dataframe.columns) + index_cols

    # -------------------------
    # Заголовки колонок
    # -------------------------
    for row in range(
        excel_start_row,
        excel_start_row + header_rows
    ):
        for col in range(
            excel_start_col,
            excel_start_col + table_width
        ):
            cell = worksheet.cell(
                row=row,
                column=col
            )

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

    # MultiIndex columns створюють
    # додатковий рядок для назви index
    extra_row = 0

    if index and header_rows > 1:
        extra_row = 1

    first_data_row = (
        excel_start_row
        + header_rows
        + extra_row
    )

    # -------------------------
    # Index таблиці
    # -------------------------
    if index:
        for row in range(
            first_data_row,
            first_data_row + len(dataframe.index)
        ):
            for col in range(
                excel_start_col,
                excel_start_col + index_cols
            ):
                cell = worksheet.cell(
                    row=row,
                    column=col
                )

                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment

def export_excel(tables):
    output_file = BytesIO()
    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        # =========================
        # BY OPERATOR
        # =========================
        write_table(
            writer,
            tables["employee_report"],
            "ByOperator",
            startrow=0,
            startcol=0,
            index=True
        )

        # =========================
        # BY LINE — основний звіт
        # =========================
        write_table(
            writer,
            tables["line_report"],
            "By Line",
            startrow=0,
            startcol=0,
            index=True
        )

        # =========================
        # BY LINE — LIBIAO
        # =========================
        write_table(
            writer,
            tables["Libiao"],
            "By Line",
            startrow=22,
            startcol=3,
            index=True,
            header=False,
            format_headers=False
        )

        # =========================
        # BY LINE — CURRENT HC
        # =========================
        write_table(
            writer,
            tables["line_hc_report"],
            "By Line",
            startrow=23,
            startcol=3,
            index=True
        )

        # =========================
        # STATISTICS
        # =========================
        write_table(
            writer,
            tables["statistics_report"],
            "Statistics",
            startrow=0,
            startcol=0,
            index=False
        )

        # =========================
        # DASHBOARD — основна таблиця
        # =========================
        write_table(
            writer,
            tables["dashboard_report"],
            "Dashboard",
            startrow=0,
            startcol=0,
            index=True
        )

        # =========================
        # DASHBOARD — RATE ZONE
        # =========================
        write_table(
            writer,
            tables["rate_zone_report"],
            "Dashboard",
            startrow=0,
            startcol=7,
            index=True
        )

        # =========================
        # DASHBOARD — HOURLY HC
        # =========================
        write_table(
            writer,
            tables["hc_final"],
            "Dashboard",
            startrow=8,
            startcol=7,
            index=True
        )



    return output_file.getvalue()
